import unittest
import tempfile
from pathlib import Path
from docx import Document
from app.rendering.engine_word import docx_to_html_string, ensure_docx_preview

class WordHtmlRenderingTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.doc_path = Path(self.temp_dir.name) / "test_doc.docx"
        doc = Document()
        doc.add_heading("Заголовок документа", level=1)
        doc.add_paragraph("Тестовый абзац с текстом для проверки HTML-рендера.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Ячейка 1"
        table.cell(0, 1).text = "Ячейка 2"
        table.cell(1, 0).text = "Ячейка 3"
        table.cell(1, 1).text = "Ячейка 4"
        doc.save(str(self.doc_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_to_html_string(self):
        html, stats = docx_to_html_string(self.doc_path)
        self.assertIn("<!doctype html>", html.lower())
        self.assertIn("launcher-sheet-zoom", html)
        self.assertGreaterEqual(stats["paragraphs"], 2)
        self.assertEqual(stats["tables"], 1)

    def test_ensure_docx_preview(self):
        cache_dir = Path(self.temp_dir.name) / "cache"
        result = ensure_docx_preview(self.doc_path, cache_dir)
        self.assertTrue((cache_dir / "preview.html").is_file())
        self.assertIn("Заголовок документа", (cache_dir / "preview.html").read_text(encoding="utf-8"))
        self.assertEqual(result["tables"], 1)
        # Test cache hit
        hit = ensure_docx_preview(self.doc_path, cache_dir)
        self.assertTrue(hit.get("cacheHit"))

if __name__ == "__main__":
    unittest.main()
